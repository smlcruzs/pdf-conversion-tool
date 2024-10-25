from flask import Flask, render_template, request, redirect, url_for, session, send_file
import pandas as pd
import os
from openpyxl import load_workbook

app = Flask(__name__)
app.config['UPLOAD_FOLDER'] = 'uploads'
app.secret_key = 'a_very_secret_key'  # Substitua por uma chave secreta complexa

@app.route('/', methods=['GET', 'POST'])
def index():
    return render_template('index.html')

@app.route('/data', methods=['POST'])
def data():
    if 'upload-file' not in request.files:
        return "No file part"
    
    file = request.files['upload-file']
    
    if file.filename == '':
        return "No selected file"
    
    if file and file.filename.endswith('.xlsx'):
        file_path = os.path.join(app.config['UPLOAD_FOLDER'], file.filename)
        file.save(file_path)
        
        # Armazenar o caminho do arquivo na sessão
        session['file_path'] = file_path
        
        return redirect(url_for('edit'))
    else:
        return "Invalid file format. Please upload an Excel file."

@app.route('/edit', methods=['GET', 'POST'])
def edit():
    file_path = session.get('file_path', None)
    if not file_path:
        return redirect(url_for('index'))
    
    if request.method == 'POST':
        try:
            # Carregar o arquivo Excel
            wb = load_workbook(file_path, data_only=False)
            ws = wb.active  # Ou use wb['NomeDaPlanilha'] se você souber o nome da planilha
            
            # Iterar sobre todas as células da planilha e aplicar lógica
            for row in ws.iter_rows():
                for cell in row:
                    # Aplicar lógica a cada célula
                    if cell.value == 'Some value':  # Substitua 'Some value' pela condição desejada
                        # Atualizar a célula ao lado (por exemplo, coluna seguinte)
                        next_cell = ws.cell(row=cell.row, column=cell.column + 1)
                        next_cell.value = 'Updated based on condition'
            
            # Salvar as alterações no arquivo Excel
            updated_file_path = os.path.join(app.config['UPLOAD_FOLDER'], 'updated_' + os.path.basename(file_path))
            wb.save(updated_file_path)
            
            # Atualizar a visualização com a planilha atualizada
            session['download_file'] = updated_file_path
            
            # Recarregar e exibir a planilha atualizada
            data = pd.read_excel(updated_file_path)
            data_dict = data.to_dict(orient='records')
            return render_template('edit.html', data=data_dict, message="File updated successfully!")
        except Exception as e:
            app.logger.error(f"Error processing data: {e}")
            return f"An error occurred: {e}"
    
    # Para métodos GET
    data = pd.read_excel(file_path)
    data_dict = data.to_dict(orient='records')
    return render_template('edit.html', data=data_dict)

@app.route('/download')
def download():
    file_path = session.get('download_file', None)
    if file_path and os.path.exists(file_path):
        return send_file(file_path, as_attachment=True, download_name=os.path.basename(file_path))
    return "File not found", 404

if __name__ == '__main__':
    app.run(debug=True)
